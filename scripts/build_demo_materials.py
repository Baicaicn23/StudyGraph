"""生成演示素材文件（PDF / DOCX / PPTX / 图片）——**全部内容自编，不引用任何教材**。

为什么要有这个脚本：演示需要「真材实料」的多格式资料（PDF 讲义的页码出处、Word 笔记、
PPT 课件、拍照板书的图片），但又**不能收编任何受版权保护的教材或网上下载的资料**——
参赛材料里有一份「来源与授权清单」，收外部内容就得逐条披露授权。所以这里的内容
全部是按公开知识自写的教学材料，只用于演示「导入 → 解析 → 检索 → 引用」这条链路。

生成到 `assets/demo_materials/`（**受版本控制**，评委拿到仓库也能复现），
再由 `zhiji.interfaces.demo`（`uv run zhiji-demo`）把它们走真实入库路径播种进知识库。

依赖是**一次性**的，用 uv 临时安装，不进项目 venv：

    HTTPS_PROXY=http://127.0.0.1:8086 uv run --with reportlab --with python-docx \\
        --with python-pptx --with pillow python scripts/build_demo_materials.py

（不需要联网时可省略 proxy 前缀。）
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "demo_materials"

# PDF 需要嵌入中文字体，否则中文会变成黑块；macOS 自带这份带 CJK 的 Unicode 字体。
CJK_FONT = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")

# ---------------------------------------------------------------- 内容（自编）

# PDF 讲义：{相对路径: (文档标题, [(页标题, [段落...]), ...])}
PDFS: dict[str, tuple[str, list[tuple[str, list[str]]]]] = {
    "高等数学-微积分/极限与连续·复习讲义.pdf": (
        "极限与连续 · 复习讲义",
        [
            (
                "本篇要点",
                [
                    "本篇按四步复习极限：先建立直观理解，再练四则运算，然后背熟两个重要极限与等价无穷小，最后处理连续性与间断点。",
                    "考试里极限几乎不会单独出题，它总是藏在导数定义、级数收敛、积分上限函数这些题的第一步。所以这一章的目标不是“会算”，而是“看到就能反应”。",
                    "复习建议：把等价无穷小表默写一遍，再把两个重要极限的变形各推一次。",
                ],
            ),
            (
                "一、极限的直观定义与记号",
                [
                    "函数 f(x) 在 x 无限接近 a 时，函数值趋于某个确定的数 L，就说 x→a 时 f(x) 的极限是 L。",
                    "关键点：极限只关心 a 点附近的行为，**不要求 f 在 a 点有定义**，也不要求 f(a) 等于 L。",
                    "记号读法：lim(x→a) f(x) = L。左极限写成 x→a⁻，右极限写成 x→a⁺。",
                    "极限存在的充要条件：左右极限都存在且相等。算分段函数在分界点处的极限时，必须两边都算。",
                ],
            ),
            (
                "二、极限的四则运算",
                [
                    "若 lim f(x) 与 lim g(x) 都存在，则：和差的极限等于极限的和差；积的极限等于极限的积；商的极限等于极限的商（要求分母的极限不为 0）。",
                    "推论：常数可以提到极限号外面；有限个函数相加相乘，极限可以逐个取。",
                    "注意“有限个”是必需的：无限项相加不能逐项取极限（这正是级数要单独讨论收敛的原因）。",
                    "遇到 0/0、∞/∞ 这类未定式，四则运算法则**不能直接用**，要先变形：约分、通分、有理化、或使用等价无穷小与洛必达法则。",
                ],
            ),
            (
                "三、两个重要极限与等价无穷小",
                [
                    "重要极限一：lim(x→0) sin x / x = 1。它的作用是**把三角函数换掉**——凡是 sin(某个趋于 0 的量) 都能用那个量替代。",
                    "重要极限二：lim(x→∞) (1 + 1/x)^x = e，等价形式 lim(t→0) (1 + t)^(1/t) = e。它处理“1 的无穷次方”型未定式。",
                    "x→0 时常用等价无穷小：sin x ~ x，tan x ~ x，arcsin x ~ x，1 − cos x ~ x²/2，ln(1 + x) ~ x，e^x − 1 ~ x。",
                    "使用铁律：等价替换只能用在**乘除**里，不能用在加减里。",
                    "反例：lim(x→0) (tan x − sin x)/x³，若把 tan x 和 sin x 都替换成 x，分子会变成 0，得出错误答案 0；正确做法是先提出 sin x 再化简，结果是 1/2。",
                ],
            ),
            (
                "四、连续性与间断点分类",
                [
                    "函数在 a 点连续，要同时满足三件事：在 a 点有定义、极限存在、且极限值等于函数值。缺一不可。",
                    "间断点分两类。第一类：左右极限都存在，但和函数值不相等（可去间断点），或左右极限存在且不相等（跳跃间断点）。第二类：至少一侧极限不存在（如无穷间断点、振荡间断点）。",
                    "易错点：把「可去间断点」说成「极限不存在」。可去间断点的极限是**存在**的，只是函数值没跟上。",
                    "可导一定连续，连续不一定可导；连续一定有极限，有极限不一定连续。背书时按这条链推，比死记四个命题可靠。",
                ],
            ),
        ],
    ),
    "线性代数/矩阵与行列式·习题课讲义.pdf": (
        "矩阵与行列式 · 习题课讲义",
        [
            (
                "本篇要点",
                [
                    "矩阵是“数排成的矩形表”，行列式是“由方阵算出来的一个数”。这两件事从定义上就不是一类对象。",
                    "本篇覆盖三件事：矩阵运算、行列式的性质与计算、可逆与秩的判定。",
                ],
            ),
            (
                "一、矩阵的运算",
                [
                    "加法要求同型（行数、列数都相同），对应位置相加。",
                    "乘法 A·B 要求 A 的列数等于 B 的行数，结果是一个 (行数 A × 列数 B) 的矩阵。",
                    "矩阵乘法**不满足交换律**：一般 AB ≠ BA。做乘法时先确认顺序。",
                    "转置的性质：(AB)ᵀ = BᵀAᵀ——注意顺序反过来。",
                ],
            ),
            (
                "二、行列式的性质与计算",
                [
                    "只有方阵才有行列式。记作 |A| 或 det A。",
                    "性质一：交换两行，行列式变号。性质二：某行乘 k，行列式也乘 k（注意与矩阵数乘的区别：矩阵数乘是每行都乘）。",
                    "性质三：两行成比例（含相等），行列式为 0。",
                    "计算策略：用初等行变换把行列式化成上三角，然后对角线相乘。三阶以上优先用这个方法，不要硬展开。",
                ],
            ),
            (
                "三、可逆与秩",
                [
                    "n 阶方阵 A 可逆，等价于 |A| ≠ 0，等价于 A 的秩等于 n，等价于 A 的行（列）向量组线性无关。",
                    "求逆的常用方法：伴随矩阵法 A⁻¹ = A*/|A|（适合二阶、三阶），或者用初等行变换 [A | E] → [E | A⁻¹]。",
                    "秩的含义：矩阵中线性无关的行（列）向量的最大个数，也就是“独立信息的条数”。",
                    "易错点：以为秩就是行数或列数。只有当行（列）全部线性无关时，秩才等于行（列）数。",
                ],
            ),
        ],
    ),
    "大学物理/牛顿运动定律·习题课讲义.pdf": (
        "牛顿运动定律 · 习题课讲义",
        [
            (
                "本篇要点",
                [
                    "本篇把三定律的表述、常见力的分析、以及解题的标准步骤串起来。重点不是背公式，而是**画对受力图**。",
                ],
            ),
            (
                "一、三定律的准确表述",
                [
                    "第一定律：物体不受外力（或所受合外力为零）时保持静止或匀速直线运动。它的价值在于**定义了惯性参考系**——牛顿定律只在这一类参考系里成立。",
                    "第二定律：F = ma。这里的 F 是合外力，a 是与力同方向的加速度。它是瞬时关系，力一变加速度立刻变。",
                    "第三定律：作用力与反作用力大小相等、方向相反、作用在**两个不同物体**上，且同时产生同时消失。",
                    "易错点：把作用力与反作用力当成一对平衡力。平衡力作用在**同一个**物体上，作用反作用力作用在**两个**物体上——这是区分二者的关键。",
                ],
            ),
            (
                "二、常见力的分析",
                [
                    "重力：mg，方向竖直向下。",
                    "支持力（弹力）：垂直于接触面，指向被支持的物体。它的大小不需要等于重力，只有在竖直静止的简单情形下才相等。",
                    "摩擦力：静摩擦力大小由平衡条件决定，最大值不超过 μsN；滑动摩擦力大小等于 μkN，方向与相对滑动方向相反。",
                    "易错点：把静摩擦力当成恒等于 μsN。静摩擦力是“需要多少给多少”，只有到了将要滑动的临界状态才等于最大值。",
                ],
            ),
            (
                "三、解题步骤与易错点",
                [
                    "标准步骤：① 选研究对象（整体法还是隔离法）；② 画受力图，只画受到的力，不画它施给别人的力；③ 建立坐标系，通常沿加速度方向建正方向；④ 按方向列牛顿第二定律的分量式；⑤ 联立求解并检验量纲。",
                    "连接体问题：先用整体法求加速度，再用隔离法求内力。顺序反了会多出未知数。",
                    "临界问题：抓住“刚要滑动”“刚要离开接触面”这些条件，此时对应摩擦力或支持力取边界值。",
                ],
            ),
        ],
    ),
    "概率论与数理统计/随机变量与分布·复习讲义.pdf": (
        "随机变量与分布 · 复习讲义",
        [
            (
                "本篇要点",
                [
                    "本篇覆盖随机变量、分布函数、常见分布、期望与方差四块。核心是把“离散用求和、连续用积分”这条主线记牢。",
                ],
            ),
            (
                "一、随机变量与分布函数",
                [
                    "随机变量是把样本空间里的每个结果映射成一个实数的函数。取值为有限或可列个的叫离散型，取值充满一个区间的叫连续型。",
                    "分布函数 F(x) = P(X ≤ x)，对离散型和连续型都适用。它是单调不减、右连续、且 x→−∞ 时趋于 0、x→+∞ 时趋于 1。",
                    "离散型用分布律 P(X = xk) = pk；连续型用概率密度 f(x)，此时 P(a < X ≤ b) 等于 f 在 [a, b] 上的积分。",
                    "易错点：连续型随机变量在单点处的概率等于 0，所以 P(a ≤ X ≤ b) 与 P(a < X < b) 的结果相同。这也说明密度函数在某点的值不是概率。",
                ],
            ),
            (
                "二、常见分布",
                [
                    "离散：0-1 分布（一次试验成功与否）、二项分布 B(n, p)（n 次独立重复试验的成功次数）、泊松分布 P(λ)（单位时间内的稀有事件次数）。",
                    "连续：均匀分布 U(a, b)（区间内等可能）、指数分布 E(λ)（等待时间，具有无记忆性）、正态分布 N(μ, σ²)（对称钟形，标准化后查标准正态表）。",
                    "二项分布在 n 很大、p 很小时可用泊松分布近似；这是“稀有事件近似”的常见考点。",
                    "正态分布的可加性：相互独立的正态变量之和仍是正态变量，均值相加、方差相加。",
                ],
            ),
            (
                "三、期望与方差",
                [
                    "期望 E(X) 刻画平均水平：离散型是 Σ xk·pk，连续型是 ∫ x·f(x) dx。",
                    "方差 D(X) = E[(X − E(X))²] = E(X²) − [E(X)]²。后一个式子算起来更快。",
                    "性质：E(aX + b) = aE(X) + b；D(aX + b) = a²D(X)（注意常数 b 不影响方差）。",
                    "易错点：以为 D(X + Y) = D(X) + D(Y) 恒成立。只有 X 与 Y **相互独立**时（更弱地，不相关时）方差才可加。",
                ],
            ),
        ],
    ),
}

# Word 笔记：{相对路径: (文档标题, [(小标题, [段落...]), ...])}
DOCXS: dict[str, tuple[str, list[tuple[str, list[str]]]]] = {
    "高等数学-微积分/导数与微分·课堂笔记.docx": (
        "导数与微分 · 课堂笔记",
        [
            (
                "导数的定义",
                [
                    "导数是函数在某一点的瞬时变化率，定义为差商的极限：f'(x0) = lim(Δx→0) [f(x0+Δx) − f(x0)] / Δx。",
                    "几何意义：曲线 y = f(x) 在点 (x0, f(x0)) 处切线的斜率。",
                    "物理意义：位移对时间的导数是速度，速度对时间的导数是加速度。",
                    "左右导数都存在且相等，才说该点可导。分段函数在分界点必须用定义分别算左右导数。",
                ],
            ),
            (
                "求导法则",
                [
                    "四则：乘积法则 (uv)' = u'v + uv'；商法则 (u/v)' = (u'v − uv')/v²。",
                    "复合函数链式法则：dy/dx = dy/du · du/dx。这是求导里用得最多的一条。",
                    "常见导数： (xⁿ)' = n·xⁿ⁻¹，(eˣ)' = eˣ，(ln x)' = 1/x，(sin x)' = cos x。",
                ],
            ),
            (
                "微分与近似",
                [
                    "微分 dy = f'(x) dx，它是函数增量的线性主部。",
                    "近似计算：当 |Δx| 很小时，f(x0 + Δx) ≈ f(x0) + f'(x0)·Δx。",
                    "易错点：把“可导”和“连续”混为一谈。可导一定连续，连续不一定可导——|x| 在 0 点连续但不可导。",
                ],
            ),
        ],
    ),
    "线性代数/特征值与对角化·复习提纲.docx": (
        "特征值与对角化 · 复习提纲",
        [
            (
                "定义与几何意义",
                [
                    "对方阵 A，若存在非零向量 v 和数 λ 使得 Av = λv，则称 λ 是 A 的特征值，v 是对应的特征向量。",
                    "几何意义：特征向量的方向在变换后保持不变，只被拉伸（或压缩）了 λ 倍。λ 为负表示反向。",
                    "求法：解特征方程 |A − λE| = 0 得到特征值，再对每个 λ 解 (A − λE)v = 0 得到特征向量。",
                ],
            ),
            (
                "对角化",
                [
                    "n 阶方阵 A 可对角化的充要条件：有 n 个线性无关的特征向量。",
                    "充分条件：n 个特征值互不相同，则一定可对角化。",
                    "重根情形：若特征值的重数等于它对应的线性无关特征向量个数（几何重数等于代数重数），仍可对角化。",
                    "对角化的用途：把 A 写成 PΛP⁻¹，于是 A^k = PΛ^kP⁻¹，求高次幂非常快。",
                ],
            ),
            (
                "实对称矩阵",
                [
                    "实对称矩阵的特征值全是实数，不同特征值对应的特征向量相互正交。",
                    "实对称矩阵一定可以正交对角化：A = QΛQᵀ，其中 Q 是正交矩阵。",
                    "易错点：把“特征向量”和“基向量”搞混。特征向量是特殊的、方向不变的向量，不一定沿坐标轴方向。",
                ],
            ),
        ],
    ),
    "大学英语四级/四级写作·常用句式与模板.docx": (
        "四级写作 · 常用句式与模板",
        [
            (
                "开头段：引出话题",
                [
                    "Nowadays, ... has become a heated topic among college students.",
                    "With the rapid development of ..., more and more attention has been paid to ...",
                    "It is widely acknowledged that ... plays an essential role in our daily life.",
                ],
            ),
            (
                "主体段：展开论述",
                [
                    "To begin with, ... What's more, ... Last but not least, ...（三段递进，结构最清晰）",
                    "There are two main reasons accounting for this phenomenon. On the one hand, ...; on the other hand, ...",
                    "Take my own experience as an example. ...",
                ],
            ),
            (
                "结尾段：总结与建议",
                [
                    "In conclusion, ... Only in this way can we ...",
                    "Therefore, it is high time that we took effective measures to ...",
                    "As far as I am concerned, ... is worth our attention.",
                ],
            ),
            (
                "提分提示",
                [
                    "字数不足时优先补“原因句”和“例子句”，不要重复题目原句凑字数。",
                    "把 very、good、many 这类初级词换成 extremely、beneficial、a considerable number of。",
                    "全文至少用一次定语从句或非谓语结构，这是阅卷时的加分观察点。",
                ],
            ),
        ],
    ),
}

# PPT 课件：{相对路径: (文档标题, [(页标题, [要点...]), ...])}
PPTXS: dict[str, tuple[str, list[tuple[str, list[str]]]]] = {
    "数据结构/二叉搜索树与平衡树·课堂讲义.pptx": (
        "二叉搜索树与平衡树 · 课堂讲义",
        [
            (
                "为什么需要二叉搜索树",
                [
                    "有序数组查找是 O(log n)，但插入删除要搬移元素，是 O(n)。",
                    "链表的插入删除是 O(1)，但查找要遍历，是 O(n)。",
                    "二叉搜索树想把两者的优点合起来：查找、插入、删除的平均代价都是 O(log n)。",
                ],
            ),
            (
                "定义与性质",
                [
                    "每个节点最多两个子节点，左子树的所有键小于根，右子树的所有键大于根。",
                    "中序遍历的结果一定是一个递增序列——这是判断一棵树是不是 BST 的最快方法。",
                ],
            ),
            (
                "查找",
                [
                    "从根开始比较：相等则命中，小于往左走，大于往右走。",
                    "时间复杂度取决于树高 h，平均 O(log n)，最坏 O(n)。",
                ],
            ),
            (
                "插入",
                [
                    "按查找的方式走到空位置，把新节点挂上去。",
                    "插入不会改变已有节点的相对结构，所以不需要旋转。",
                ],
            ),
            (
                "删除（三种情况）",
                [
                    "叶子节点：直接删除。",
                    "只有一个子节点：用子节点顶替它的位置。",
                    "有两个子节点：用右子树的最小节点（中序后继）或左子树的最大节点替换，再递归删除那个替代节点。",
                ],
            ),
            (
                "退化问题",
                [
                    "如果按有序序列依次插入，BST 会退化成一条链，树高变成 n。",
                    "此时所有操作都退化成 O(n)，和链表没有区别。",
                    "这正是「平衡树」要解决的问题。",
                ],
            ),
            (
                "平衡树如何补救",
                [
                    "AVL 树：要求任意节点左右子树高度差不超过 1，插入删除后用旋转维持。查找最稳定，但旋转频繁。",
                    "红黑树：用颜色约束近似平衡，旋转次数少，插入删除更快；标准库里的 map / set 多用它。",
                ],
            ),
            (
                "易错点小结",
                [
                    "把「平均 O(log n)」当成「一定 O(log n)」——最坏情况是 O(n)。",
                    "删除有两个子节点的节点时，忘了递归删除用来替换的那个节点，导致出现重复键。",
                    "只用「左孩子小于根」判断 BST，忽略了整棵左子树都必须小于根。",
                ],
            ),
        ],
    ),
    "计算机网络/TCP连接管理·课堂讲义.pptx": (
        "TCP 连接管理 · 课堂讲义",
        [
            (
                "TCP 的两个核心特征",
                [
                    "面向连接：通信前必须先建立连接，通信后要释放。",
                    "可靠传输：确认、重传、序号、滑动窗口共同保证不丢不乱。",
                ],
            ),
            (
                "三次握手",
                [
                    "第一次：客户端发 SYN，带自己的初始序号 seq = x。",
                    "第二次：服务端回 SYN + ACK，自己的序号 seq = y，确认号 ack = x + 1。",
                    "第三次：客户端回 ACK，确认号 ack = y + 1，连接建立。",
                ],
            ),
            (
                "为什么是三次而不是两次",
                [
                    "两次无法确认客户端的接收能力。",
                    "更实际的原因：防止历史失效的连接请求突然到达服务端，造成服务端白白建立连接。",
                ],
            ),
            (
                "四次挥手",
                [
                    "主动方发 FIN；被动方先回 ACK（此时半关闭，被动方可能还有数据要发）。",
                    "被动方数据发完后发自己的 FIN；主动方回 ACK，进入 TIME_WAIT。",
                    "因为被动方的 ACK 和 FIN 不能合并，所以是四次而不是三次。",
                ],
            ),
            (
                "TIME_WAIT 为什么要等 2MSL",
                [
                    "保证最后那个 ACK 能到达对端；丢了的话对端会重传 FIN，还有机会补发。",
                    "让本次连接的旧报文在网络中彻底消失，避免污染下一个使用同样端口的新连接。",
                ],
            ),
            (
                "易错点小结",
                [
                    "把三次握手说成“确认双方都能收发”——准确说法是确认双方收发能力，并同步初始序号。",
                    "忘记 TIME_WAIT 是主动关闭方进入的状态，不是被动方。",
                ],
            ),
        ],
    ),
}

# 图片（用于演示「拍照课件 / 板书 → 视觉转写」）
IMAGES: dict[str, tuple[str, list[str], tuple[str, str]]] = {
    # 相对路径: (画面标题, 行内容, (背景色, 文字色))
    "演示用图片/板书-极限的运算法则.png": (
        "极限的运算法则",
        [
            "1. lim(f ± g) = lim f ± lim g",
            "2. lim(f · g) = lim f · lim g",
            "3. lim(f / g) = lim f / lim g  (lim g ≠ 0)",
            "",
            "两个重要极限：",
            "  lim(x→0) sinx / x = 1",
            "  lim(x→∞) (1 + 1/x)^x = e",
            "",
            "x→0 等价无穷小：sinx ~ x, 1-cosx ~ x²/2",
        ],
        ("#1f3a34", "#f2f7f4"),
    ),
    "演示用图片/课件截图-二叉树遍历.png": (
        "二叉树的三种遍历",
        [
            "前序：根 → 左 → 右",
            "中序：左 → 根 → 右",
            "后序：左 → 右 → 根",
            "",
            "性质：二叉搜索树的中序遍历是递增序列",
            "查找 / 插入 / 删除 平均 O(log n)，最坏 O(n)",
        ],
        ("#ffffff", "#1f2937"),
    ),
}


def _render_pdf(path: Path, title: str, pages: list[tuple[str, list[str]]]) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    if not CJK_FONT.exists():
        raise SystemExit(f"找不到中文字体：{CJK_FONT}（PDF 会渲染成黑块，已中止）")
    pdfmetrics.registerFont(TTFont("CJK", str(CJK_FONT)))

    h1 = ParagraphStyle("h1", fontName="CJK", fontSize=18, leading=26, spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName="CJK", fontSize=13, leading=20, spaceBefore=8, spaceAfter=6)
    body = ParagraphStyle("body", fontName="CJK", fontSize=10.5, leading=17, spaceAfter=5)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, title=title,
        leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
    )
    flow = [Paragraph(title, h1), Spacer(1, 6)]
    for index, (page_title, paragraphs) in enumerate(pages):
        if index:
            flow.append(PageBreak())
        flow.append(Paragraph(page_title, h2))
        for text in paragraphs:
            flow.append(Paragraph(text, body))
    doc.build(flow)


def _render_docx(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(title, level=0)
    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for text in paragraphs:
            doc.add_paragraph(text)
    doc.save(str(path))


def _render_pptx(path: Path, title: str, slides: list[tuple[str, list[str]]]) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    # 用默认模板的「标题和内容」版式（layout 1）逐页铺开
    layout = prs.slide_layouts[1]
    for index, (slide_title, bullets) in enumerate(slides):
        slide = prs.slides.add_slide(layout)
        # 首页标题带上课件名，方便在阅读器/检索结果里认出这是哪份课件
        slide.shapes.title.text = f"{title}｜{slide_title}" if index == 0 else slide_title
        body = slide.placeholders[1].text_frame
        body.text = bullets[0]
        for line in bullets[1:]:
            body.add_paragraph().text = line
        for paragraph in body.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(18)
    prs.save(str(path))


def _render_image(path: Path, title: str, lines: list[str], colors: tuple[str, str]) -> None:
    from PIL import Image, ImageDraw, ImageFont

    background, foreground = colors
    width, height = 1000, 620
    image = Image.new("RGB", (width, height), background)
    draw = ImageDraw.Draw(image)
    font_title = ImageFont.truetype(str(CJK_FONT), 44)
    font_body = ImageFont.truetype(str(CJK_FONT), 28)
    draw.text((60, 48), title, font=font_title, fill=foreground)
    draw.line((60, 112, width - 60, 112), fill=foreground, width=2)
    y = 150
    for line in lines:
        if line:
            draw.text((60, y), line, font=font_body, fill=foreground)
        y += 46
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(str(path))


def main() -> int:
    created: list[Path] = []
    for relative, (title, pages) in PDFS.items():
        target = OUT / relative
        _render_pdf(target, title, pages)
        created.append(target)
    for relative, (title, sections) in DOCXS.items():
        target = OUT / relative
        _render_docx(target, title, sections)
        created.append(target)
    for relative, (title, slides) in PPTXS.items():
        target = OUT / relative
        _render_pptx(target, title, slides)
        created.append(target)
    for relative, (title, lines, colors) in IMAGES.items():
        target = OUT / relative
        _render_image(target, title, lines, colors)
        created.append(target)

    print(f"已生成 {len(created)} 个演示素材 → {OUT.relative_to(ROOT)}/")
    for item in created:
        print(f"  {item.relative_to(OUT)}  ({item.stat().st_size / 1024:.1f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
